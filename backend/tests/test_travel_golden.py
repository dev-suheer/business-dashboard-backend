"""Golden-file tests for travel, against the sample workbooks.

Every figure here is checked against the client's own worked reconciliation in
claim 13 (its ``Analysis`` and ``Sheet1`` tabs), or against the claim's own
summary block. Where ours differs from theirs on purpose, the test says so.
"""

import zipfile
from datetime import date
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from app.recon.categories.travel.export import REASON_LABELS, build_report
from app.recon.categories.travel.models import (
    BookingMatchState,
    BookingRecon,
    RefundKind,
    TravelRateConfig,
    TravelResult,
    UnmatchedReason,
    Vertical,
)
from app.recon.categories.travel.pipeline import TravelRequest, run
from app.recon.exceptions import ExceptionCode, Severity
from app.recon.inputs import Upload
from tests.conftest import (
    TRAVEL_CLAIMS,
    TRAVEL_PORTAL,
    TRAVEL_WALLET,
    write_finance_report,
)

RATES = TravelRateConfig(
    hotel_rate=Decimal("0.05"),
    flight_rate=Decimal("0.03"),
    vat_rate=Decimal("0.15"),
    base_tolerance=Decimal("0.05"),
    commission_tolerance=Decimal("0.01"),
    rate_tolerance=Decimal("0.0005"),
)


def reconcile(claim: str, *, wallet: bool = True) -> TravelResult:
    return run(
        TravelRequest(
            portal=Upload(role="portal", filename="portal.xlsx", path=TRAVEL_PORTAL),
            vendor=Upload(
                role="vendor",
                filename=TRAVEL_CLAIMS[claim].name,
                path=TRAVEL_CLAIMS[claim],
            ),
            wallet=(
                Upload(role="wallet", filename="wallet.xlsx", path=TRAVEL_WALLET)
                if wallet
                else None
            ),
            rates=RATES,
            period_start=None,
            period_end=None,
            max_rows=1_000_000,
        )
    )


@pytest.fixture(scope="module")
def result() -> TravelResult:
    """Claim 13 — the one week the client also reconciled by hand."""
    return reconcile("13")


def booking(result: TravelResult, booking_id: str) -> BookingRecon:
    return next(b for b in result.bookings if b.booking_id == booking_id)


def test_the_whole_export_is_read(result: TravelResult) -> None:
    assert result.summary.portal_rows == 206
    assert result.summary.portal_bookings == 45
    assert result.period_start == date(2026, 7, 26)
    assert result.period_end == date(2026, 8, 22)
    assert result.vendor_name == "Almosafer"


def test_the_claim_is_read_whole(result: TravelResult) -> None:
    assert result.summary.vendor_lines == 8
    assert result.summary.vendor_bookings == 7
    assert result.summary.bookings_matched == 7
    # Every booking this claim bills has a portal row.
    assert result.summary.bookings_vendor_only == 0


def test_the_claim_totals_agree_with_the_claims_own_summary(
    result: TravelResult,
) -> None:
    """Claim 13's Sheet2 reads 6985.16 / 255.19765 / 6729.96235."""
    assert result.summary.claim_base == Decimal("6985.16")
    assert result.summary.claim_commission == Decimal("255.20")
    assert result.summary.claim_net_payable == Decimal("6729.96")
    # Every booking this claim bills matched, so the matched side is all of it.
    assert result.summary.vendor_base == result.summary.claim_base


def test_the_hotel_base_difference_matches_the_clients_sheet(
    result: TravelResult,
) -> None:
    """Their Sheet1: 1161.70 against 1162.76, a difference of -1.06."""
    hotel = booking(result, "H260819107345500")
    assert hotel.portal_base == Decimal("1161.70")
    assert hotel.vendor_base == Decimal("1162.76")
    assert hotel.base_diff == Decimal("-1.06")
    assert result.summary.base_difference == Decimal("-1.06")


def test_the_settled_figures_match_the_clients_sheet(result: TravelResult) -> None:
    """Their 'Actual As Vendor' 66.79775 and 'Net Payable' 1094.90225, which
    this reports quantized once, as every money figure is."""
    hotel = booking(result, "H260819107345500")
    assert hotel.locality.value == "international"
    assert hotel.contract_rate == Decimal("0.0575")
    assert hotel.commission_as_vendor == Decimal("66.80")
    assert hotel.net_payable_as_vendor == Decimal("1094.90")
    # Their 'Difference': -0.06095, at the contract rate on both sides.
    assert hotel.commission_diff_as_vendor == Decimal("-0.06")
    assert result.summary.commission_difference_at_contract == Decimal("-0.06")
    assert result.summary.net_payable_difference_at_contract == Decimal("-1.00")


def test_the_vat_uplift_is_measured(result: TravelResult) -> None:
    """We charged 5%; the claim bills 5.75%. The difference is uncollected.

    Totalled over the matched bookings only — the four international ones,
    4.31 + 5.96 + 8.71 + 3.39 — so each booking is counted on the claim that
    bills it rather than on every week's run.
    """
    hotel = booking(result, "H260819107345500")
    assert hotel.portal_commission == Decimal("58.09")
    assert hotel.sg_expected_commission == Decimal("58.09")
    assert hotel.vat_uplift_shortfall == Decimal("8.71")
    assert result.summary.vat_uplift_shortfall == Decimal("22.37")
    shortfall = next(
        e
        for e in result.exceptions
        if e.code is ExceptionCode.SG_COMMISSION_SHORT_OF_VENDOR_RATE
    )
    assert shortfall.amount == Decimal("22.37")


def test_a_local_booking_is_billed_at_our_own_rate(result: TravelResult) -> None:
    local = booking(result, "A26082185665500")
    assert local.locality.value == "local"
    assert local.contract_rate == Decimal("0.03")
    assert local.portal_commission == local.vendor_commission == Decimal("28.05")
    assert local.vat_uplift_shortfall == Decimal(0)


def test_a_multi_leg_booking_is_compared_once(result: TravelResult) -> None:
    """Two invoice lines, 481.38 + 477.41, against one portal row of 958.79."""
    multi = booking(result, "A260813120395500")
    assert multi.vendor_lines == 2
    assert multi.vendor_base == Decimal("958.79")
    assert multi.portal_base == Decimal("958.79")
    assert multi.base_diff == Decimal(0)
    assert "MULTI_LEG" in multi.flags


def test_bookings_billed_on_another_week_are_not_breaks(result: TravelResult) -> None:
    """The export spans a month; this claim covers one week of it."""
    codes = result.summary.exception_counts
    assert codes["PRIOR_PERIOD"] == 34
    assert codes["BOOKING_MISSING_IN_VENDOR"] == 4


def test_the_payment_capture_defect_is_found(result: TravelResult) -> None:
    """SG order 198 collected 13.00 of a 276.25 total."""
    split = next(
        e for e in result.exceptions if e.code is ExceptionCode.PAYMENT_SPLIT_MISMATCH
    )
    assert split.amount == Decimal("263.25")
    assert len(result.portal_rows_flagged) == 3


def test_every_payment_method_is_counted(result: TravelResult) -> None:
    """Every booked row is counted; rows with no booking ID are left out."""
    by_label = {m.label: m for m in result.payment_methods.by_label}
    assert set(by_label) == {
        "Card",
        "Tamara",
        "Apple Pay",
        "Wallet",
        "Apple Pay + Wallet",
        "Card + Wallet",
        "Tamara + Wallet",
    }
    assert sum(m.rows for m in result.payment_methods.by_label) == 45
    by_instrument = {m.label: m for m in result.payment_methods.by_instrument}
    # The wallet column totals what actually moved on wallets, combined
    # payments included.
    assert by_instrument["Wallet"].order_total == Decimal("7832.93")


def test_the_wallet_history_accounts_for_every_wallet_payment(
    result: TravelResult,
) -> None:
    wallet = result.wallet
    assert wallet is not None
    assert (wallet.window_start, wallet.window_end) == (
        date(2026, 8, 1),
        date(2026, 8, 15),
    )
    # Booked rows only: the ten wallet-paid rows with no booking ID are left out.
    assert wallet.rows_checked == 24
    assert wallet.matched == 17
    assert wallet.amount_mismatch == 0
    assert wallet.missing == 0
    # The rest fall outside the history's own two-week window.
    assert wallet.out_of_window == 7
    # Leaving a row out of the check does not make its movement portal-less.
    assert len(wallet.unmatched_history) == 1


def test_no_wallet_file_means_no_wallet_block() -> None:
    assert reconcile("13", wallet=False).wallet is None


def test_the_overcharge_and_its_credit_are_both_found() -> None:
    """Claim 10 bills A260629226265500 at 4.4724%; claim 12 credits it back."""
    overcharge = next(
        e
        for e in reconcile("10", wallet=False).exceptions
        if e.code is ExceptionCode.VENDOR_COMMISSION_OVERCHARGE
    )
    assert overcharge.order_no == "A260629226265500"
    assert overcharge.amount == Decimal("66.82")

    credited = reconcile("12", wallet=False)
    assert credited.summary.vendor_adjustments == Decimal("-66.82")
    applied = next(
        e for e in credited.exceptions if e.code is ExceptionCode.ADJUSTMENT_APPLIED
    )
    assert applied.order_no == "A260629226265500"


def test_a_partial_refund_reports_what_was_retained() -> None:
    """Claim 10 credits A26070648655500 less an 824 fee and a 30 change fee."""
    result = reconcile("10", wallet=False)
    refund = booking(result, "A26070648655500")
    assert refund.refund is RefundKind.PARTIAL
    assert refund.match_state is BookingMatchState.VENDOR_ONLY
    assert refund.vendor_cancel_fee == Decimal("824.00")
    assert refund.vendor_change_fee == Decimal("30.00")
    assert result.summary.retained_fees == Decimal("854.00")


def test_a_refund_the_portal_never_shows_is_taken_from_the_claim() -> None:
    """Claim 12 credits H260810138915500 in full; the portal still shows 850."""
    result = reconcile("12", wallet=False)
    refunded = booking(result, "H260810138915500")
    assert refunded.refund is RefundKind.FULL
    assert refunded.portal_base == Decimal("850.00")
    assert refunded.vendor_base == Decimal(0)
    # The vendor reversed its commission; we did not reverse ours.
    assert refunded.vendor_commission == Decimal(0)
    assert refunded.portal_commission == Decimal("42.50")
    kept = [
        e
        for e in result.exceptions
        if e.code is ExceptionCode.COMMISSION_NOT_REVERSED_ON_REFUND
    ]
    assert [e.amount for e in kept] == [Decimal("42.50")]


def test_a_refunded_booking_is_not_counted_as_a_base_difference() -> None:
    """Its two sides are meant to disagree, so it belongs on the refunds sheet."""
    result = reconcile("12", wallet=False)
    assert result.summary.base_difference == Decimal("-3.86")


def test_the_claims_own_ledger_is_carried_through(result: TravelResult) -> None:
    assert [p.claim_period for p in result.claim_periods] == [
        "Claim From 21 till 27 Jun 2026",
        "claim From 09 till 15 AUG 2026",
        "claim From 16 till 22 AUG 2026",
    ]


def test_uploading_the_same_claim_twice_would_not_double_it() -> None:
    """Every line of every claim is distinct under the deduplication key."""
    for name in TRAVEL_CLAIMS:
        result = reconcile(name, wallet=False)
        assert ExceptionCode.DUPLICATE_VENDOR_LINE not in {
            e.code for e in result.exceptions
        }


def test_the_exception_list_leads_with_what_needs_action(
    result: TravelResult,
) -> None:
    severities = [e.severity for e in result.exceptions]
    assert severities == sorted(
        severities, key=lambda s: ["high", "medium", "info"].index(str(s))
    )
    assert severities[0] is Severity.HIGH


SOURCES = {"portal": TRAVEL_PORTAL, "vendor": TRAVEL_CLAIMS["13"]}


def test_the_report_carries_every_sheet(result: TravelResult, tmp_path: Path) -> None:
    path = build_report(result, tmp_path / "travel.xlsx", SOURCES)
    book = openpyxl.load_workbook(path, read_only=True)
    assert book.sheetnames == [
        "Dashboard",
        "Orders",
        "Exceptions",
        "Wallet",
        "B2C Invoices ERP",
        "Portal",
        "Vendor",
    ]
    book.close()


def test_the_report_omits_the_wallet_sheet_when_no_wallet_was_uploaded(
    tmp_path: Path,
) -> None:
    path = build_report(reconcile("13", wallet=False), tmp_path / "travel.xlsx")
    with zipfile.ZipFile(path) as book:
        sheets = book.read("xl/workbook.xml").decode()
    assert 'name="Wallet"' not in sheets
    assert 'name="B2C Invoices ERP"' not in sheets


def test_hotel_and_flight_commission_split_with_its_ex_vat_column(
    result: TravelResult,
) -> None:
    s = result.summary
    assert s.hotel_commission + s.flight_commission == s.commission_as_vendor
    # Matched bookings only: 66.80 / 1.15 and 188.33 / 1.15.
    assert s.hotel_commission_ex_vat == Decimal("58.09")
    assert s.flight_commission_ex_vat == Decimal("163.77")


def test_every_vertical_is_resolved(result: TravelResult) -> None:
    assert {b.vertical for b in result.bookings} == {Vertical.HOTEL, Vertical.FLIGHT}


def test_only_bookings_both_sides_agree_on_are_matched(result: TravelResult) -> None:
    """Claim 13 bills seven of the export's 45 bookings; the rest are explained."""
    s = result.summary
    assert s.bookings_matched == 7
    assert s.bookings_unmatched == 38
    assert s.unmatched_by_reason == {"cancelled": 3, "missing_in_vendor": 35}
    unmatched = [b for b in result.bookings if b.unmatched_reason is not None]
    assert len(unmatched) == 38
    assert all(b.explanation for b in unmatched)
    # The totals are taken from the matched bookings alone.
    matched = [b for b in result.bookings if b.unmatched_reason is None]
    assert s.portal_base == sum((b.portal_base for b in matched), Decimal(0))


def test_a_booking_billed_on_another_week_says_so(result: TravelResult) -> None:
    early = booking(result, "A260726150195500")
    assert early.unmatched_reason is UnmatchedReason.MISSING_IN_VENDOR
    assert "before this claim's week" in early.explanation
    assert "expected on another claim" in early.explanation


def test_a_refund_matched_by_id_is_explained_not_matched() -> None:
    """Both sides carry H260810138915500, but the claim credits it in full."""
    result = reconcile("12", wallet=False)
    refunded = booking(result, "H260810138915500")
    assert refunded.match_state is BookingMatchState.MATCHED
    assert refunded.unmatched_reason is UnmatchedReason.REFUNDED
    assert "full refund" in refunded.explanation
    cancelled = booking(result, "H260809145945500")
    assert cancelled.unmatched_reason is UnmatchedReason.CANCELLED
    assert "no credit note" in cancelled.explanation
    assert result.summary.unmatched_by_reason == {
        "refunded": 1,
        "cancelled": 2,
        "missing_in_vendor": 33,
        "missing_in_portal": 1,
    }


def test_a_vendor_only_booking_is_missing_from_the_portal() -> None:
    result = reconcile("10", wallet=False)
    prior = booking(result, "A260629226265500")
    assert prior.unmatched_reason is UnmatchedReason.MISSING_IN_PORTAL
    assert "before the portal export" in prior.explanation
    credit = booking(result, "A26070648655500")
    assert credit.unmatched_reason is UnmatchedReason.REFUNDED
    assert "854.00 in fees retained" in credit.explanation


def test_rows_without_a_booking_id_are_kept_but_not_calculated(
    result: TravelResult,
) -> None:
    assert result.summary.portal_rows_without_booking_id == 161
    unbooked = [r for r in result.references.portal_rows if r.booking_id is None]
    assert len(unbooked) == 161
    assert result.wallet is not None
    assert all(check.booking_id for check in result.wallet.checks)
    # Their exceptions are still raised: two expired orders kept commission.
    assert result.summary.exception_counts["COMMISSION_ON_UNCOMPLETED_ORDER"] == 2


@pytest.mark.parametrize("claim", ["10", "11", "12", "13"])
def test_every_portal_row_is_reported_exactly_once(claim: str) -> None:
    """A matched booking's own row, or a reference: never both, never neither."""
    result = reconcile(claim, wallet=False)
    references = {(r.sheet, r.row) for r in result.references.portal_rows}
    assert len(references) == len(result.references.portal_rows)
    assert result.summary.bookings_matched + len(references) == (
        result.summary.portal_rows
    )


@pytest.mark.parametrize("claim", ["12", "13"])
def test_the_report_splits_matched_from_unmatched(claim: str, tmp_path: Path) -> None:
    """Matched bookings are totalled; every other booking is listed beneath."""
    result = reconcile(claim)
    path = build_report(result, tmp_path / "travel.xlsx")
    book = openpyxl.load_workbook(path, read_only=True)
    rows = list(book["Orders"].iter_rows(values_only=True))
    headers = list(rows[0])
    match = headers.index("MATCH ID'S")
    verdict = headers.index("Verdict")

    s = result.summary
    total = next(i for i, r in enumerate(rows) if str(r[0] or "").startswith("Total"))
    matched = rows[1:total]
    assert len(matched) == s.bookings_matched
    assert all(r[match] == "TRUE" for r in matched)
    title = next(
        i
        for i, r in enumerate(rows)
        if str(r[0] or "").startswith("Bookings not in the totals")
    )
    unmatched = [r for r in rows[title + 1 :] if any(v is not None for v in r)]
    assert len(unmatched) == s.bookings_unmatched
    reasons = set(REASON_LABELS.values())
    assert all(r[verdict] in reasons for r in unmatched)
    book.close()


def test_the_orders_sheet_mirrors_the_clients_analysis(
    result: TravelResult, tmp_path: Path
) -> None:
    path = build_report(result, tmp_path / "travel.xlsx", SOURCES)
    book = openpyxl.load_workbook(path)
    sheet = book["Orders"]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    for header in (
        "Partner Order ID",
        "MATCH ID'S",
        "Com Rate Check",
        "Vendor Rate Check",
        "Actual As Vendor",
        "Total Invoice Amount",
        "Verdict",
    ):
        assert header in headers
    # Portal green, checks light green, claim green, as on food's sheet.
    assert sheet.cell(1, 1).fill.fgColor.rgb == "FF375623"
    totals = 2 + result.summary.bookings_matched
    order_total = headers.index("Order Total") + 1
    column = sheet.cell(2, order_total).column_letter
    assert sheet.cell(totals, order_total).value == f"=SUM({column}2:{column}8)"
    hotel = next(
        row
        for row in range(2, totals)
        if sheet.cell(row, headers.index("Partner Order ID") + 1).value
        == "H260819107345500"
    )
    assert sheet.cell(hotel, headers.index("BASE DIFFERENCE") + 1).value == -1.06
    assert sheet.cell(hotel, headers.index("Verdict") + 1).value == "Base differs"
    labels = {
        sheet.cell(row, 1).value: row
        for row in range(totals, sheet.max_row + 1)
        if sheet.cell(row, 1).value
    }
    base = labels["Base Amount"]
    assert sheet.cell(base, 6).value == f"=D{base}-E{base}"
    assert sheet.cell(base - 1, 5).value == "ALMOSAFER"
    book.close()


def test_the_dashboard_reads_the_matched_bookings(
    result: TravelResult, tmp_path: Path
) -> None:
    path = build_report(result, tmp_path / "travel.xlsx", SOURCES)
    book = openpyxl.load_workbook(path)
    dashboard = book["Dashboard"]
    labels = {
        dashboard.cell(row, 1).value: row
        for row in range(1, dashboard.max_row + 1)
        if dashboard.cell(row, 1).value
    }
    header = labels["SG TRAVEL ORDER Amount"]
    values = {
        dashboard.cell(header, col).value: dashboard.cell(header + 1, col).value
        for col in range(3, 18)
    }
    s = result.summary
    assert values["Order"] == s.bookings_matched == 7
    assert Decimal(str(values["Net Payable"])) == s.net_payable_as_vendor
    assert Decimal(str(values["Booking Amount"])) == s.portal_base
    breakdown = labels["Breakdown By Most Used"]
    assert dashboard.cell(breakdown, 4).value == "Revenue By Provider"
    total = breakdown + 1
    assert dashboard.cell(total, 1).value == "TOTAL"
    # A booking billed by two airlines is one order, under each airline.
    assert dashboard.cell(total, 2).value == 7
    providers = range(total + 1, dashboard.max_row + 1)
    assert sum(dashboard.cell(row, 2).value for row in providers) == 8
    assert Decimal(str(dashboard.cell(total, 10).value)) == s.net_payable_as_vendor
    book.close()


def test_b2c_invoices_bill_each_matched_booking_once(
    result: TravelResult, tmp_path: Path
) -> None:
    path = build_report(result, tmp_path / "travel.xlsx", SOURCES)
    book = openpyxl.load_workbook(path)
    invoices = book["B2C Invoices ERP"]
    orders = book["Orders"]
    lines = list(invoices.iter_rows(min_row=2, values_only=True))
    # The matched bookings only: refunded, cancelled and one-sided ones are not.
    assert len(lines) == result.summary.bookings_matched == 7
    assert len({line[1] for line in lines}) == len(lines)
    by_booking = {line[1]: line for line in lines}
    assert "A260726150195500" not in by_booking
    hotel = by_booking["H260819107345500"]
    assert hotel[0] == "Almosafer"
    assert hotel[2] == "حسن فقيه"
    assert hotel[4] == "Hotel Booking"
    assert hotel[7] == "Domestic Sales (15%)"
    cell = hotel[6]
    assert cell.startswith("=Orders!") and cell.endswith("/1.15")
    reference = cell[len("=Orders!") : -len("/1.15")]
    # The formula points at that booking's Commission Charged on Orders.
    column = "".join(ch for ch in reference if ch.isalpha())
    row = int("".join(ch for ch in reference if ch.isdigit()))
    assert orders[f"{column}1"].value == "Commission Charged"
    assert orders[f"D{row}"].value == "H260819107345500"
    assert orders[f"{column}{row}"].value == 58.09
    assert by_booking["A26082185665500"][4] == "Flight Booking"
    book.close()


def _formula(cell: object) -> str:
    """openpyxl returns a dynamic-array formula as an object, not a string."""
    value = getattr(cell, "value", None)
    text = getattr(value, "text", None)
    return text if isinstance(text, str) else str(value)


def test_a_finance_report_adds_cashback_lookups(
    result: TravelResult, tmp_path: Path
) -> None:
    finance = write_finance_report(tmp_path / "finance.xlsx", [121], "hotel")
    path = build_report(
        result, tmp_path / "travel.xlsx", {**SOURCES, "finance": finance}
    )
    book = openpyxl.load_workbook(path)
    assert book.sheetnames[-1] == "Finance Report"
    sheet = book["Orders"]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    assert headers[headers.index("Verdict") + 1 :] == [
        "Cash Back Reward",
        "Vertical Name",
        "Reward Code",
        "Reward Type",
        "Outcome",
    ]
    assert headers[2] == "SG Order ID"
    reward_col = headers.index("Cash Back Reward") + 1
    reward = _formula(sheet.cell(2, reward_col))
    assert reward.startswith('=IF(C2="","NA",_xlfn.IFNA(_xlfn.XLOOKUP(')
    assert "'Finance Report'!C:C" in reward
    assert "'Finance Report'!F:F" in reward
    totals = 2 + result.summary.bookings_matched
    assert str(sheet.cell(totals, 1).value).startswith("Total")
    # The rewards are not totalled.
    assert sheet.cell(totals, reward_col).value is None
    # Bookings listed under the totals are looked up too.
    title = next(
        row
        for row in range(totals, sheet.max_row + 1)
        if str(sheet.cell(row, 1).value or "").startswith("Bookings not in")
    )
    below = _formula(sheet.cell(title + 1, reward_col))
    assert below.startswith(f'=IF(C{title + 1}="","NA",')
    copied = book["Finance Report"]
    assert copied.cell(15, 3).value == 121
    book.close()
