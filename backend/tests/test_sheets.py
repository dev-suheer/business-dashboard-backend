"""The report helpers every category shares."""

from pathlib import Path

import xlsxwriter
from openpyxl import Workbook, load_workbook

from app.recon.categories.food import parser
from app.recon.models import SheetManifest
from app.recon.sheets import letter, portal_records, report_formats, source_sheets

PORTAL_HEADERS = (
    "SG Order ID",
    "Partner Order ID",
    "Vertical",
    "Order Total",
    "SG Delivery Expense",
    "Commission Charged",
    "Net Payable Amount to Partner",
    "Order Status",
    "Total Amount",
    "Restaurant Name",
)


def _upload(path: Path, sheets: dict[str, list[tuple[object, ...]]]) -> Path:
    book = Workbook()
    book.remove(book.active)
    for name, rows in sheets.items():
        sheet = book.create_sheet(name)
        sheet.append(PORTAL_HEADERS)
        for values in rows:
            sheet.append(values)
            # Forced to text: a cell that only looks like a formula.
            cell = sheet.cell(sheet.max_row, len(PORTAL_HEADERS))
            cell.data_type = "s"
    book.save(path)
    return path


def test_uploaded_text_is_copied_as_text_never_as_a_formula(tmp_path: Path) -> None:
    upload = _upload(
        tmp_path / "portal.xlsx",
        {"Sheet": [(1, 99, "Food", 10, 2, 1, 9, "Delivered", 10, "=1+1")]},
    )
    report = tmp_path / "report.xlsx"
    book = xlsxwriter.Workbook(str(report))
    source_sheets(
        book,
        report_formats(book),
        [
            SheetManifest(
                role="portal", filename="portal.xlsx", sheets=["Sheet"], rows=1
            )
        ],
        {"portal": upload},
        (("portal", "Portal", parser.PORTAL_SCHEMA, "portal_header"),),
    )
    book.close()

    copied = load_workbook(report)["Portal"].cell(2, len(PORTAL_HEADERS))
    assert copied.value == "=1+1"
    assert copied.data_type == "s"


def test_stacked_sheets_are_numbered_the_way_the_copy_writes_them(
    tmp_path: Path,
) -> None:
    """Each stacked sheet gets a name row above it and a blank row after it."""
    upload = _upload(
        tmp_path / "portal.xlsx",
        {
            "Hotel": [(1, "H1", "Food", 10, 2, 1, 9, "Delivered", 10, "a")],
            "Flight": [(2, "A1", "Food", 10, 3, 1, 9, "Delivered", 10, "b")],
        },
    )
    records = list(portal_records(upload, ["Hotel", "Flight"], parser.PORTAL_SCHEMA))
    assert [(r.sheet, r.ordinal, r.excel_row) for r in records] == [
        ("Hotel", 1, 3),
        ("Flight", 1, 7),
    ]
    assert records[1].cell("sgdeliveryexpense") == "Portal!E7"
    assert letter(0) == "A" and letter(25) == "Z" and letter(26) == "AA"
