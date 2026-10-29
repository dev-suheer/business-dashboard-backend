"""The food endpoint: uploads, formats and the errors it must give back."""

from decimal import Decimal
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.core.config import settings
from tests.conftest import (
    FOOD_PORTAL,
    FOOD_VENDOR,
    TRAVEL_PORTAL,
    TRAVEL_WALLET,
    VENDOR_STATEMENT,
)

FOOD = f"{settings.API_V1_STR}/reconcile/food"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _files() -> dict[str, tuple[str, bytes, str]]:
    return {
        "portal": ("portal.xlsx", FOOD_PORTAL.read_bytes(), XLSX),
        "vendor": ("vendor.xlsx", FOOD_VENDOR.read_bytes(), XLSX),
    }


def test_end_to_end_json(auth_client: TestClient) -> None:
    response = auth_client.post(FOOD, files=_files())
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["category"] == "food"
    assert body["vendor_name"] == "The Chefz"
    assert body["summary"]["settled_orders"] == 297
    assert Decimal(str(body["summary"]["settled_due"])) == Decimal("20748.07")
    detected = {item["role"]: item["sheets"] for item in body["detected"]}
    assert detected["portal"] == ["PORTAL REPORT"]
    assert detected["vendor"] == ["VENDOR REPORT"]
    assert len(body["orders"]) == 306


def test_end_to_end_xlsx(auth_client: TestClient) -> None:
    response = auth_client.post(FOOD, files=_files(), params={"format": "xlsx"})
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == XLSX
    assert (
        "reconciliation-food-2026-09-01-to-2026-09-06.xlsx"
        in (response.headers["content-disposition"])
    )
    book = load_workbook(BytesIO(response.content), read_only=True)
    assert book.sheetnames == [
        "Dashboard",
        "Orders",
        "Exceptions",
        "B2C Invoices ERP",
        "Portal",
        "Vendor",
    ]
    book.close()


def test_the_wallet_file_is_optional(auth_client: TestClient) -> None:
    assert auth_client.post(FOOD, files=_files()).json()["wallet"] is None
    files = _files()
    files["wallet"] = ("wallet_history.xlsx", TRAVEL_WALLET.read_bytes(), XLSX)
    response = auth_client.post(FOOD, files=files)
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["wallet"]["unkeyed"] == 188
    assert Decimal(str(body["summary"]["settled_due"])) == Decimal("20748.07")
    detected = {item["role"]: item["sheets"] for item in body["detected"]}
    assert detected["wallet"] == ["Sheet1"]


def test_a_missing_file_is_a_validation_error(auth_client: TestClient) -> None:
    assert auth_client.post(FOOD).status_code == 422


def test_a_travel_portal_has_no_food_rows(auth_client: TestClient) -> None:
    files = _files()
    files["portal"] = ("portal.xlsx", TRAVEL_PORTAL.read_bytes(), XLSX)
    response = auth_client.post(FOOD, files=files)
    assert response.status_code == 422
    assert "no food orders" in response.json()["detail"]


def test_an_account_statement_is_rejected_by_name(auth_client: TestClient) -> None:
    files = _files()
    files["vendor"] = ("vendor_statement.xlsx", VENDOR_STATEMENT.read_bytes(), XLSX)
    response = auth_client.post(FOOD, files=files)
    assert response.status_code == 422
    assert "food vendor report" in response.json()["detail"]


def test_a_damaged_finance_file_is_a_422_naming_the_field(
    auth_client: TestClient,
) -> None:
    files = _files()
    files["finance"] = ("finance.xlsx", b"not a workbook", XLSX)
    response = auth_client.post(FOOD, files=files)
    assert response.status_code == 422
    assert "'finance' is not a readable Excel workbook" in response.json()["detail"]


def test_a_non_excel_upload_is_refused(auth_client: TestClient) -> None:
    files = _files()
    files["portal"] = ("portal.csv", b"not a workbook", "text/csv")
    response = auth_client.post(FOOD, files=files)
    assert response.status_code == 415


def _finance_bytes() -> bytes:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    headers = [
        "Spin Event ID",
        "User ID",
        "Order ID",
        "Reward Amount (SAR)",
        "Reward Code",
        "Reward Type",
        "Outcome",
        "Vertical Name",
    ]
    for col, header in enumerate(headers, start=1):
        sheet.cell(1, col, header)
    sheet.cell(2, 3, 59709)
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def test_a_finance_report_is_optional_and_copied(auth_client: TestClient) -> None:
    files = _files()
    files["finance"] = ("finance.xlsx", _finance_bytes(), XLSX)
    response = auth_client.post(FOOD, files=files, params={"format": "xlsx"})
    assert response.status_code == 200, response.text
    book = load_workbook(BytesIO(response.content), read_only=True)
    assert book.sheetnames[-1] == "Finance Report"
    book.close()


def test_a_finance_file_without_the_spin_table_is_rejected(
    auth_client: TestClient,
) -> None:
    files = _files()
    files["finance"] = ("finance.xlsx", FOOD_VENDOR.read_bytes(), XLSX)
    response = auth_client.post(FOOD, files=files)
    assert response.status_code == 422
    assert "finance" in response.json()["detail"]


def test_the_endpoint_needs_a_session(client: TestClient) -> None:
    assert client.post(FOOD, files=_files()).status_code == 401
