"""The eShop workbook follows the food one, and the wallet history is checked."""

from decimal import Decimal
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from app.core.config import settings
from app.recon.engine import ReconRequest, Upload, reconcile
from app.recon.exceptions import ExceptionCode
from app.recon.export import build_report
from app.recon.models import RateConfig, ReconResult
from app.recon.wallet import WalletState
from tests.conftest import TRAVEL_WALLET, VENDOR_STATEMENT, WORKINGS

SOURCES = {"portal": WORKINGS, "vendor": VENDOR_STATEMENT}


def _reconcile(*, wallet: bool) -> ReconResult:
    rates = RateConfig(
        standard_rate=settings.STANDARD_RATE,
        promo_rate=settings.PROMO_RATE,
        promo_cutoff=settings.PROMO_CUTOFF,
        line_tolerance=settings.LINE_TOLERANCE,
        order_tolerance=settings.ORDER_TOLERANCE,
    )
    return reconcile(
        ReconRequest(
            category="eshop",
            portal=Upload("portal", WORKINGS.name, WORKINGS),
            promo=Upload("promo", WORKINGS.name, WORKINGS),
            vendor=Upload("vendor", VENDOR_STATEMENT.name, VENDOR_STATEMENT),
            settlements=None,
            rates=rates,
            period_start=None,
            period_end=None,
            max_rows=1_000_000,
            wallet=(
                Upload("wallet", TRAVEL_WALLET.name, TRAVEL_WALLET) if wallet else None
            ),
        )
    )


@pytest.fixture(scope="module")
def result() -> ReconResult:
    return _reconcile(wallet=True)


def _labels(sheet: Any, column: int = 1) -> dict[Any, int]:
    return {
        sheet.cell(row, column).value: row
        for row in range(1, sheet.max_row + 1)
        if sheet.cell(row, column).value
    }


def test_the_wallet_does_not_move_the_payable(result: ReconResult) -> None:
    assert result.summary.net_payable == Decimal("13876.31")
    assert _reconcile(wallet=False).wallet is None


def test_an_order_is_found_in_the_wallet_by_its_number(result: ReconResult) -> None:
    """SG-SH-08986 is SG Order ID 8986: 241.99 from the wallet on 15 Aug."""
    wallet = result.wallet
    assert wallet is not None
    checks = {check.booking_id: check for check in wallet.checks}
    found = checks["SG-SH-08986"]
    assert found.sg_order_id == "8986"
    assert found.state is WalletState.MATCHED
    assert found.history_paid == Decimal("241.99")
    # The payment column repeats on every line: never summed.
    assert checks["SG-SH-08194"].portal_wallet_amount == Decimal("467.60")
    # The history covers 1 to 15 August even though its eShop movements are
    # all on the 15th, so an order on the 11th is missing, not out of window.
    assert checks["SG-SH-08624"].state is WalletState.MISSING
    assert [event.sg_order_id for event in wallet.unmatched_history] == ["8920"]
    codes = {exc.code for exc in result.exceptions}
    assert ExceptionCode.WALLET_HISTORY_UNMATCHED in codes
    assert result.summary.exception_counts["WALLET_ENTRY_MISSING"] == wallet.missing


def test_the_report_has_the_food_sheets(result: ReconResult, tmp_path: Path) -> None:
    path = build_report(result, tmp_path / "eshop.xlsx", SOURCES)
    book = load_workbook(path, read_only=True)
    assert book.sheetnames == [
        "Dashboard",
        "Orders",
        "Exceptions",
        "Wallet",
        "Cashback",
        "Portal",
        "Vendor",
    ]
    book.close()


def test_the_dashboard_carries_the_headline(
    result: ReconResult, tmp_path: Path
) -> None:
    path = build_report(result, tmp_path / "eshop.xlsx", SOURCES)
    book = load_workbook(path)
    dashboard = book["Dashboard"]
    labels = _labels(dashboard)
    header = labels["SG ESHOP ORDER Amount"]
    values = {
        dashboard.cell(header, col).value: dashboard.cell(header + 1, col).value
        for col in range(3, 18)
    }
    included = [order for order in result.orders if order.included_in_payable]
    assert values["Order"] == len(included) == 34
    assert Decimal(str(values["Net Payable"])) == Decimal("13876.31")
    assert Decimal(str(values["Contract Commission"])) == (
        result.summary.commission_earned
    )
    shipping = labels["Shipping"]
    assert dashboard.cell(shipping + 1, 3).value == 510
    breakdown = labels["Breakdown By Most Used"]
    assert dashboard.cell(breakdown, 4).value == "Revenue By Brand"
    total = breakdown + 1
    assert dashboard.cell(total, 1).value == "TOTAL"
    assert dashboard.cell(total, 2).value == 34
    assert Decimal(str(dashboard.cell(total, 9).value)) == (
        result.summary.commission_earned
    )
    # Multi-brand orders count under each brand, so the rows outnumber orders.
    brands = range(total + 1, dashboard.max_row + 1)
    assert sum(dashboard.cell(row, 2).value for row in brands) > 34
    book.close()


def test_the_orders_sheet_carries_the_settlement_bridge(
    result: ReconResult, tmp_path: Path
) -> None:
    path = build_report(result, tmp_path / "eshop.xlsx", SOURCES)
    book = load_workbook(path)
    sheet = book["Orders"]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    for header in ("Order Number", "MATCH ID'S", "Invoice Amount", "Verdict"):
        assert header in headers
    totals = 2 + 34
    assert str(sheet.cell(totals, 1).value).startswith("Total, 34")
    bridge = totals + 2
    titles = [sheet.cell(bridge, col).value for col in range(1, 11)]
    assert titles[2] == "Net payable"
    assert sheet.cell(bridge + 1, 3).value == 13876.31
    assert sheet.cell(bridge + 1, 4).value == 17424.75
    verdicts = {
        sheet.cell(row, headers.index("Order Number") + 1).value: sheet.cell(
            row, headers.index("Verdict") + 1
        ).value
        for row in range(2, sheet.max_row + 1)
    }
    assert verdicts["SG-SH-03344"] == "Clawback"
    assert verdicts["SG-SH-07435"] == "Return adjusted"
    labels = _labels(sheet)
    assert "Orders not in the payable, not in the totals" in labels
    statement = labels["Statement rows with no portal order, not in the totals"]
    prior_or_later = [
        sheet.cell(row, headers.index("Verdict") + 1).value
        for row in range(statement + 1, sheet.max_row + 1)
    ]
    assert len(prior_or_later) == 21
    assert set(prior_or_later) == {"Prior period", "Later period"}
    book.close()


def test_the_cashback_bills_sg_shipping(result: ReconResult, tmp_path: Path) -> None:
    path = build_report(result, tmp_path / "eshop.xlsx", SOURCES)
    book = load_workbook(path)
    cashback = book["Cashback"]
    portal = book["Portal"]
    lines = list(cashback.iter_rows(min_row=2, values_only=True))
    # Every order with SG-owned shipping, which the export puts on its first line.
    assert len(lines) == 37
    first = lines[0]
    assert first[0] == "RM"
    assert first[1] == "SG-SH-03344"
    assert first[2] is None
    assert first[4] == "eShop Delivery"
    assert first[5] == 1
    assert first[6] == "=Portal!Y2/1.15"
    assert first[7] == "Domestic Sales (15%)"
    assert portal["Y1"].value == "Shipping Cost"
    assert portal["Y2"].value == 15
    book.close()


def test_exceptions_carry_the_order_status(result: ReconResult, tmp_path: Path) -> None:
    path = build_report(result, tmp_path / "eshop.xlsx", SOURCES)
    book = load_workbook(path)
    sheet = book["Exceptions"]
    headers = [sheet.cell(1, col).value for col in range(1, 7)]
    assert headers == ["Severity", "Code", "Order", "Status", "Amount", "Detail"]
    statuses = {
        sheet.cell(row, 3).value: sheet.cell(row, 4).value
        for row in range(2, sheet.max_row + 1)
    }
    assert statuses["SG-SH-05422"] == "RTO"
    book.close()
