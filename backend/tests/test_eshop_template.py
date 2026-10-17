"""The vendor statement template (app/recon/categories/eshop/template.py).

A template is only worth handing out if the parser reads what vendors type into
it, and never mistakes its instructions or worked example for data. These tests
build it, fill it the way a vendor would, and run it back through the parser and
the whole reconciliation.
"""

from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from app.recon.categories.eshop import template
from app.recon.categories.eshop.parser import (
    LEDGER_SCHEMA,
    classify_doc_type,
    read_ledger,
)
from app.recon.engine import ReconRequest, Upload, reconcile
from app.recon.excel import HEADER_SCAN_LIMIT, detect_sheets, find_header, open_workbook
from app.recon.exceptions import ExceptionCode
from app.recon.matcher import resolve_credit_notes
from app.recon.models import LedgerDocType, LedgerEntry, RateConfig, ReconResult
from tests.conftest import VENDOR_STATEMENT, WORKINGS

D = Decimal
MAX_ROWS = 1_000_000


@pytest.fixture
def blank(tmp_path: Path) -> Path:
    return template.build_vendor_template(tmp_path / "template.xlsx")


def _fill(blank: Path, rows: list[tuple[Any, ...]], into: Path) -> Path:
    """Type rows into the statement from the first data row down.

    Each row is (date, type, doc no, order no, description, debit, credit,
    balance). openpyxl keeps no formula results, so a balance given here is
    written over the formula, standing in for the value Excel caches on save.
    """
    workbook = load_workbook(blank)
    sheet = workbook[template.STATEMENT_SHEET]
    for offset, values in enumerate(rows):
        for col, value in enumerate(values, start=1):
            if value is not None:
                sheet.cell(
                    row=template.FIRST_DATA_ROW + offset, column=col, value=value
                )
    workbook.save(into)
    return into


def _read(path: Path) -> list[LedgerEntry]:
    with open_workbook(path) as workbook:
        return list(
            read_ledger(
                workbook, detect_sheets(workbook, LEDGER_SCHEMA), max_rows=MAX_ROWS
            )
        )


def _reconcile(vendor: Path, rates: RateConfig) -> ReconResult:
    return reconcile(
        ReconRequest(
            category="eshop",
            portal=Upload("portal", WORKINGS.name, WORKINGS),
            promo=Upload("promo", WORKINGS.name, WORKINGS),
            vendor=Upload("vendor", vendor.name, vendor),
            settlements=None,
            rates=rates,
            period_start=None,
            period_end=None,
            max_rows=MAX_ROWS,
        )
    )


# --- shape ------------------------------------------------------------------


def test_only_the_statement_sheet_is_read_as_a_ledger(blank: Path) -> None:
    # The instructions list every Arabic header and end in a worked example; the
    # parser reads every ledger-shaped sheet, so neither may look like one.
    with open_workbook(blank) as workbook:
        assert detect_sheets(workbook, LEDGER_SCHEMA) == [template.STATEMENT_SHEET]
        match = find_header(workbook[template.STATEMENT_SHEET], LEDGER_SCHEMA)
    assert match is not None
    assert match.row == template.HEADER_ROW <= HEADER_SCAN_LIMIT
    # Everything but the converter's truncated reference, which the order
    # number column replaces.
    assert set(match.columns) == {
        "doc_date",
        "doc_type",
        "doc_no",
        "order_no",
        "raw_desc",
        "debit",
        "credit",
        "balance",
    }


def test_a_blank_template_holds_no_transactions(blank: Path) -> None:
    # Every entry row carries formulas, cached blank until a spreadsheet
    # recalculates; none of them may surface as a transaction.
    assert _read(blank) == []


def test_every_dropdown_value_types_its_row() -> None:
    assert [classify_doc_type(value) for value in template.DOC_TYPES] == [
        LedgerDocType.SALES_INVOICE,
        LedgerDocType.SALES_RETURN,
        LedgerDocType.JOURNAL_ENTRY,
        LedgerDocType.JOURNAL_ENTRY,
    ]


# --- filled in --------------------------------------------------------------


def test_filled_rows_take_their_order_number_from_its_own_column(
    blank: Path, tmp_path: Path
) -> None:
    filled = _fill(
        blank,
        [
            (date(2026, 8, 3), template.SALES_INVOICE, "756954", "SG-SH-05785",
             "Order SG-SH-05785", 425.01, None, None),
            # The description names no order, as on the vendor's own returns.
            (date(2026, 8, 16), template.SALES_RETURN, "36814", "sg-sh-06118",
             "العميل لم يتم استلام طلب", None, 168, None),
            (date(2026, 8, 20), template.COMMISSION, "9931", None,
             "Commission for the July cycle", None, 1048.43, None),
            (date(2026, 8, 20), template.PAYMENT, "9932", None, None, None, 5000, None),
        ],
        tmp_path / "filled.xlsx",
    )  # fmt: skip

    entries = _read(filled)

    assert [(e.doc_type, e.order_no, e.debit, e.credit) for e in entries] == [
        (LedgerDocType.SALES_INVOICE, "SG-SH-05785", D("425.01"), None),
        (LedgerDocType.SALES_RETURN, "SG-SH-06118", None, D("168.00")),
        (LedgerDocType.JOURNAL_ENTRY, None, None, D("1048.43")),
        (LedgerDocType.JOURNAL_ENTRY, None, None, D("5000.00")),
    ]
    assert entries[0].doc_date == date(2026, 8, 3)


def test_the_sample_cycle_reconciles_identically_through_the_template(
    blank: Path, tmp_path: Path, rates: RateConfig
) -> None:
    """The converter's statement, retyped into the template, changes nothing.

    A vendor filling the template writes each return's full order number, which
    the converter's statement never carried; the engine's own prefix-and-amount
    resolution supplies it here. The template has no reference column at all,
    so a return only matches through that order-number column.

    The one figure that moves is the closing balance. The template's running
    balance is the sum of the amounts typed in, 17,424.81; the vendor's own
    statement prints 17,424.75, because its system books fractions of a cent
    behind the two-decimal amounts it shows (fourteen rows drift by 0.01).
    """
    baseline = _reconcile(VENDOR_STATEMENT, rates)
    entries = _read(VENDOR_STATEMENT)
    resolve_credit_notes(
        entries,
        {order.order_no: order.portal_gmv for order in baseline.orders},
        rates.order_tolerance,
    )

    rows = []
    # Oldest first, as the instructions ask; the balance runs as the template's
    # formula would compute it.
    balance = D(0)
    for entry in sorted(entries, key=lambda e: e.doc_date or date.min):
        balance += (entry.debit or 0) - (entry.credit or 0)
        rows.append(_template_row(entry, balance))
    result = _reconcile(_fill(blank, rows, tmp_path / "sample.xlsx"), rates)

    moved = {"vendor_closing_balance", "unexplained_difference"}
    assert result.summary.model_dump(exclude=moved) == (
        baseline.summary.model_dump(exclude=moved)
    )
    assert baseline.summary.vendor_closing_balance == D("17424.75")
    assert result.summary.vendor_closing_balance == D("17424.81")
    assert result.summary.ledger_entries == 64
    assert result.summary.orders_matched == 37
    assert result.summary.orders_exact == 33
    assert result.summary.net_payable == D("13876.31")
    assert not [
        e for e in result.exceptions if e.code is ExceptionCode.AMBIGUOUS_CREDIT_NOTE
    ]
    assert [(o.order_no, o.vendor_amount, o.flags) for o in result.orders] == [
        (o.order_no, o.vendor_amount, o.flags) for o in baseline.orders
    ]


def _template_row(entry: LedgerEntry, balance: Decimal) -> tuple[Any, ...]:
    if entry.doc_type is LedgerDocType.SALES_INVOICE:
        doc_type = template.SALES_INVOICE
    elif entry.doc_type is LedgerDocType.SALES_RETURN:
        doc_type = template.SALES_RETURN
    elif entry.raw_desc and "عمولة" in entry.raw_desc:
        doc_type = template.COMMISSION
    else:
        doc_type = template.PAYMENT
    return (
        entry.doc_date,
        doc_type,
        entry.doc_no,
        entry.order_no,
        entry.raw_desc,
        float(entry.debit) if entry.debit else None,
        float(entry.credit) if entry.credit else None,
        float(balance),
    )
