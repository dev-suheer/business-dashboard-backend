"""The Orders sheet matches the client recon layout and its summary block."""

from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook

from app.recon.categories.food.export import build_report
from app.recon.categories.food.models import FoodPortalRow, FoodResult, FoodVendorRow
from app.recon.categories.food.pipeline import FoodRequest, run
from app.recon.categories.food.rules import reconcile, summarise
from app.recon.exceptions import ExceptionCode
from app.recon.inputs import Upload
from tests.conftest import FOOD_PORTAL, FOOD_VENDOR, TRAVEL_WALLET


def _formula(cell: object) -> str:
    """openpyxl returns a dynamic-array formula as an object, not a string."""
    value = getattr(cell, "value", None)
    text = getattr(value, "text", None)
    return text if isinstance(text, str) else str(value)


def _cents(value: object) -> Decimal:
    return Decimal(str(value)).quantize(Decimal("0.01"))


def test_the_orders_sheet_has_the_client_totals_and_summary(tmp_path: Path) -> None:
    result = run(
        FoodRequest(
            portal=Upload("portal", FOOD_PORTAL.name, FOOD_PORTAL),
            vendor=Upload("vendor", FOOD_VENDOR.name, FOOD_VENDOR),
            max_rows=10_000,
        )
    )
    path = tmp_path / "food.xlsx"
    build_report(
        result,
        path,
        {"portal": FOOD_PORTAL, "vendor": FOOD_VENDOR},
    )

    book = load_workbook(path)
    assert book.sheetnames == [
        "Dashboard",
        "Orders",
        "Exceptions",
        "B2C Invoices ERP",
        "Portal",
        "Vendor",
    ]
    sheet = book["Orders"]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    assert "ORDER AMOUNT DIFFERENCE" in headers
    assert "Additional Del Charges SG BEAR" in headers
    assert "Additional Delivery CHarge CHEFZ BEAR" in headers
    assert "Customer Delivery Charges" in headers
    order_col = headers.index("Partner Order ID") + 1
    order_ids = {
        sheet.cell(row, order_col).value for row in range(2, sheet.max_row + 1)
    }
    assert "291072289" not in order_ids
    # Vendor rows with no portal order stay on Exceptions, out of these totals.
    assert "289926686" not in order_ids
    assert "Portal orders with no vendor row, not in the totals" not in {
        sheet.cell(row, 1).value for row in range(1, sheet.max_row + 1)
    }

    match_col = headers.index("MATCH ID'S") + 1
    assert sheet.cell(1, match_col).fill.fgColor.rgb == "FFC6EFCE"

    # 298 matched orders. Vendor-only and cancelled rows are not on this sheet.
    last_data = 299
    totals_row = 300
    sg_col = headers.index("Additional Del Charges SG BEAR") + 1
    chefz_col = headers.index("Additional Delivery CHarge CHEFZ BEAR") + 1
    customer_col = headers.index("Customer Delivery Charges") + 1
    fee_col = headers.index("RATE% CHECK SG-VNDR SOA") + 1

    def column_sum(column: int) -> Decimal:
        return sum(
            (
                Decimal(str(sheet.cell(row, column).value or 0))
                for row in range(2, last_data + 1)
            ),
            Decimal(0),
        )

    assert _cents(column_sum(sg_col)) == Decimal("1589.20")
    assert _cents(column_sum(chefz_col)) == Decimal("1278.80")
    assert _cents(column_sum(customer_col)) == Decimal("2148.00")
    assert sheet.cell(totals_row, fee_col).number_format == "#,##0.00"
    assert sheet.cell(2, fee_col).number_format == "#,##0.00"

    vendor_row = None
    for row in range(totals_row, sheet.max_row + 1):
        if sheet.cell(row, 1).value == "Vendor FOOD ORDER Amount":
            vendor_row = row
            break
    assert vendor_row is not None
    assert sheet.cell(vendor_row - 1, 1).value == "SG FOOD ORDER Amount"
    header_values = [sheet.cell(vendor_row, col).value for col in range(1, 20)]
    assert "Food Amount" in header_values
    assert "Delivery" in header_values
    assert "ADD DC" in header_values
    assert "Total Order Amount" in header_values
    assert "T. Remaining" in header_values
    # Labels and values sit on one horizontal block, not a vertical list.
    assert header_values.index("Food Amount") > 1
    assert sheet.cell(vendor_row + 1, 4).value.startswith("=")

    texts = [
        cell.value
        for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row)
        for cell in row
        if isinstance(cell.value, str)
    ]
    assert all("\u2014" not in value for value in texts)

    dashboard = book["Dashboard"]
    labels = {
        dashboard.cell(row, 1).value: row
        for row in range(1, dashboard.max_row + 1)
        if dashboard.cell(row, 1).value
    }
    assert labels["SG FOOD ORDER Amount"]
    assert labels["Cash"]
    header_row = labels["SG FOOD ORDER Amount"]
    headers = [dashboard.cell(header_row, col).value for col in range(3, 18)]
    assert headers[0] == "Order"
    assert "Net Payable" in headers
    assert "T. Remaining" in headers
    # 297 settled + 1 held pair + 4 vendor-only orders. Portal-only stays out.
    assert dashboard.cell(header_row + 1, 3).value == 302
    assert _cents(dashboard.cell(header_row + 1, 5).value) == Decimal("2179.00")
    assert _cents(dashboard.cell(header_row + 1, 6).value) == Decimal("1611.10")

    breakdown = labels["Breakdown By Most Used"]
    assert [dashboard.cell(breakdown, col).value for col in range(1, 11)] == [
        "Breakdown By Most Used",
        "Order Count",
        "Discount Count",
        "Revenue By Restaurant",
        "Applepay",
        "Tamara",
        "Card",
        "Wallet",
        "Commission",
        "Total Delivery",
    ]
    assert dashboard.cell(breakdown + 1, 1).value == "TOTAL"
    assert dashboard.cell(breakdown + 1, 2).value == 302
    assert (
        dashboard.cell(breakdown + 2, 2).value >= dashboard.cell(breakdown + 3, 2).value
    )

    invoices = book["B2C Invoices ERP"]
    assert [invoices.cell(1, col).value for col in range(1, 9)] == [
        "Partner",
        "Invoice Partner Display Name",
        "Contact Name",
        "Invoice/Bill Date",
        "Invoice lines/Product",
        "Invoice lines/Quantity",
        "Invoice lines/Unit Price",
        "Invoice lines/Taxes",
    ]
    invoice_ids = {
        str(invoices.cell(row, 2).value): row for row in range(2, invoices.max_row + 1)
    }
    # Portal-only and vendor-only orders are not matched.
    assert "291073584" not in invoice_ids
    assert "289926686" not in invoice_ids
    # The amount break is matched and delivered, so it is billed.
    billed = invoice_ids["290238794"]
    assert invoices.cell(billed, 1).value == "The Chefz"
    assert invoices.cell(billed, 5).value == "Food Delivery"
    assert invoices.cell(billed, 6).value == 1
    assert invoices.cell(billed, 8).value == "Domestic Sales (15%)"
    order_headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    delivery_col = order_headers.index("SG Delivery Expense") + 1
    order_rows = {
        str(sheet.cell(row, order_col).value): row
        for row in range(2, sheet.max_row + 1)
    }
    delivery_letter = ""
    number = delivery_col
    while number:
        number, remainder = divmod(number - 1, 26)
        delivery_letter = chr(65 + remainder) + delivery_letter
    assert invoices.cell(billed, 7).value == (
        f"=Orders!{delivery_letter}{order_rows['290238794']}/1.15"
    )
    assert invoices.cell(invoice_ids["291067212"], 3).value == "Ali AlZamanan"

    portal = book["Portal"]
    assert [portal.cell(1, col).value for col in range(1, 8)] == [
        "Order Date",
        "Order Time",
        "Customer",
        "Customer",
        "Customer",
        "SG Order ID",
        "Partner Order ID",
    ]
    assert portal.cell(1, 1).fill.fgColor.rgb == "FF375623"
    assert portal.max_row == 303
    assert portal.cell(2, 7).value == 291073584

    vendor = book["Vendor"]
    # The title row sits above the headers, and its labels are not in column A.
    assert vendor.cell(1, 3).value == "Common Field"
    assert vendor.cell(2, 1).value == "requestedDeliveryDate"
    assert vendor.cell(2, 3).value == "order_id"
    assert vendor.cell(2, 1).fill.fgColor.rgb == "FFA9D08E"
    assert vendor.max_row == 304
    assert vendor.cell(3, 3).value == 289926686

    exceptions = book["Exceptions"]
    headers = [exceptions.cell(1, col).value for col in range(1, 7)]
    assert headers == ["Severity", "Code", "Order", "Status", "Amount", "Detail"]
    by_order = {
        exceptions.cell(row, 3).value: exceptions.cell(row, 4).value
        for row in range(2, exceptions.max_row + 1)
    }
    assert by_order["290238794"] == "Delivered"
    assert by_order["289926686"] == "Delivered"
    assert by_order["291072289"] == "Delivered"
    book.close()


def _portal(order_id: str, status: str) -> FoodPortalRow:
    return FoodPortalRow(
        source="portal.xlsx",
        sheet="PORTAL REPORT",
        row=2,
        order_id=order_id,
        order_date=date(2026, 9, 2),
        vertical="Food",
        order_status=status,
        total_amount=Decimal(41),
        commission_charged=Decimal("1.95"),
    )


def _vendor(order_id: str, status: str) -> FoodVendorRow:
    return FoodVendorRow(
        source="vendor.xlsx",
        sheet="VENDOR REPORT",
        row=2,
        order_id=order_id,
        order_status=status,
        total=Decimal(41),
        shoppinggate_fees=Decimal("1.95"),
        original_delivery_fee=Decimal(17),
        sg_delivery_charge=Decimal("6.90"),
        chef_delivery_charge=Decimal("8.10"),
    )


def test_exception_status_names_each_side_when_they_differ(tmp_path: Path) -> None:
    portal_rows = [_portal("mixed", "Delivered"), _portal("cancelled", "Cancelled")]
    vendor_rows = [_vendor("mixed", "Cancelled"), _vendor("cancelled", "Cancelled")]
    orders, exceptions = reconcile(portal_rows, vendor_rows)
    result = FoodResult(
        orders=orders,
        exceptions=exceptions,
        summary=summarise(portal_rows, vendor_rows, orders, exceptions),
    )
    path = tmp_path / "status.xlsx"
    build_report(result, path)

    book = load_workbook(path)
    sheet = book["Exceptions"]
    found = {
        (sheet.cell(row, 3).value, sheet.cell(row, 2).value): sheet.cell(row, 4).value
        for row in range(2, sheet.max_row + 1)
    }
    assert (
        found[("mixed", "STATUS_MISMATCH")] == "Portal: Delivered / Vendor: Cancelled"
    )
    assert found[("cancelled", "STATUS_NOT_DELIVERED")] == "Cancelled"
    book.close()


def test_a_wallet_history_adds_the_wallet_sheet(tmp_path: Path) -> None:
    """The food sample leaves SG Order ID blank, so nothing can be looked up."""
    result = run(
        FoodRequest(
            portal=Upload("portal", FOOD_PORTAL.name, FOOD_PORTAL),
            vendor=Upload("vendor", FOOD_VENDOR.name, FOOD_VENDOR),
            max_rows=10_000,
            wallet=Upload("wallet", TRAVEL_WALLET.name, TRAVEL_WALLET),
        )
    )
    # The settled figures do not move.
    assert result.summary.settled_orders == 297
    assert result.summary.settled_due == Decimal("20748.07")
    wallet = result.wallet
    assert wallet is not None
    assert wallet.rows_checked == 0
    assert wallet.unkeyed == 188
    # The history's five food payments have no portal row to stand on.
    assert len(wallet.unmatched_history) == 5
    assert {event.vertical for event in wallet.unmatched_history} == {"food"}
    codes = [exc.code for exc in result.exceptions]
    assert codes.count(ExceptionCode.WALLET_ORDER_ID_MISSING) == 1
    assert codes.count(ExceptionCode.WALLET_HISTORY_UNMATCHED) == 5

    path = tmp_path / "food.xlsx"
    build_report(result, path, {"portal": FOOD_PORTAL, "vendor": FOOD_VENDOR})
    book = load_workbook(path, read_only=True)
    assert book.sheetnames == [
        "Dashboard",
        "Orders",
        "Exceptions",
        "Wallet",
        "B2C Invoices ERP",
        "Portal",
        "Vendor",
    ]
    book.close()


def _finance_report(path: Path, order_id: int) -> None:
    """A spin report with a summary block above the table, plus a decoy Orders sheet."""
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "CASHBACK"
    sheet["A1"] = "Total spins"
    sheet["B1"] = 1
    headers = [
        "Spin Event ID",
        "User ID",
        "Order ID",
        "Order Type",
        "Order Amount (SAR)",
        "Reward Amount (SAR)",
        "Reward Code",
        "Reward Type",
        "Outcome",
        "Spin Earned At (KSA)",
        "Reward Earned At (KSA)",
        "AOV Tier ID",
        "AOV Tier Name",
        "Vertical ID",
        "Vertical Name",
    ]
    for col, header in enumerate(headers, start=1):
        sheet.cell(14, col, header)
    sheet.cell(15, 1, 61167)
    sheet.cell(15, 3, order_id)
    sheet.cell(15, 6, 4.13)
    sheet.cell(15, 7, "PERCENT_10")
    sheet.cell(15, 8, 0.1)
    sheet.cell(15, 9, "won")
    sheet.cell(15, 15, "food")
    decoy = book.create_sheet("Orders")
    decoy["A1"] = "SG Order ID"
    decoy["B1"] = "Reward Amount (SAR)"
    decoy["C1"] = "Partner Order ID"
    book.save(path)


def test_a_finance_report_adds_cashback_lookups(tmp_path: Path) -> None:
    result = run(
        FoodRequest(
            portal=Upload("portal", FOOD_PORTAL.name, FOOD_PORTAL),
            vendor=Upload("vendor", FOOD_VENDOR.name, FOOD_VENDOR),
            max_rows=10_000,
        )
    )
    finance = tmp_path / "finance.xlsx"
    _finance_report(finance, 59709)
    path = tmp_path / "food.xlsx"
    build_report(
        result,
        path,
        {"portal": FOOD_PORTAL, "vendor": FOOD_VENDOR, "finance": finance},
    )

    book = load_workbook(path)
    assert book.sheetnames[-1] == "Finance Report"
    sheet = book["Orders"]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    assert headers[-5:] == [
        "Cash Back Reward",
        "Vertical Name",
        "Reward Code",
        "Reward Type",
        "Outcome",
    ]
    assert headers[headers.index("Verdict") + 1] == "Cash Back Reward"
    reward = _formula(sheet.cell(2, headers.index("Cash Back Reward") + 1))
    assert reward.startswith('=IF(C2="","NA",IFNA(')
    assert "XLOOKUP(IFERROR(VALUE(C2),C2)" in reward
    assert "'Finance Report'!C:C" in reward
    assert "'Finance Report'!F:F" in reward
    vertical = _formula(sheet.cell(2, headers.index("Vertical Name") + 1))
    assert "'Finance Report'!O:O" in vertical
    code = _formula(sheet.cell(2, headers.index("Reward Code") + 1))
    assert "'Finance Report'!G:G" in code
    copied = book["Finance Report"]
    assert copied.cell(14, 3).value == "Order ID"
    assert copied.cell(15, 3).value == 59709
    book.close()
