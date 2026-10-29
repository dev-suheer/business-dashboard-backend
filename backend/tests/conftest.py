import json
from collections.abc import Iterator
from decimal import Decimal
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.core import credentials as credentials_module
from app.core.config import settings
from app.core.security import get_password_hash
from app.main import app
from app.recon.categories.travel.models import TravelRateConfig
from app.recon.models import RateConfig

FIXTURES = Path(__file__).parent / "fixtures"
WORKINGS = FIXTURES / "workings.xlsx"
VENDOR_STATEMENT = FIXTURES / "vendor_statement.xlsx"

TRAVEL_FIXTURES = FIXTURES / "travel"
TRAVEL_PORTAL = TRAVEL_FIXTURES / "portal.xlsx"
TRAVEL_WALLET = TRAVEL_FIXTURES / "wallet_history.xlsx"
# One weekly claim per file, as they arrive. Claim 13 keeps every tab it came
# with, including the client's own Analysis sheet, because that is what proves
# the schemas do not mistake a worked reconciliation for an input.
TRAVEL_CLAIMS = {
    n: TRAVEL_FIXTURES / f"claim_{n}.xlsx" for n in ("10", "11", "12", "13")
}

FOOD_FIXTURES = FIXTURES / "food"
FOOD_PORTAL = FOOD_FIXTURES / "portal.xlsx"
FOOD_VENDOR = FOOD_FIXTURES / "vendor.xlsx"

TEST_USERNAME = "tester"
TEST_PASSWORD = "correct horse battery"


@pytest.fixture(autouse=True)
def credential_file(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    """Point the app at a throwaway credential file.

    Autouse so no test can accidentally read or rewrite the real
    ``credentials.json``.
    """
    path = tmp_path / "credentials.json"
    path.write_text(
        json.dumps(
            {
                "username": TEST_USERNAME,
                "password_hash": get_password_hash(TEST_PASSWORD),
            }
        )
    )
    monkeypatch.setattr(credentials_module, "credentials_path", lambda: path)
    credentials_module.clear_credential_cache()
    yield path
    credentials_module.clear_credential_cache()


@pytest.fixture
def client() -> Iterator[TestClient]:
    with TestClient(app) as test_client:
        yield test_client


@pytest.fixture
def auth_client(client: TestClient) -> TestClient:
    """A client holding a live session cookie."""
    response = client.post(
        f"{settings.API_V1_STR}/login/access-token",
        data={"username": TEST_USERNAME, "password": TEST_PASSWORD},
    )
    assert response.status_code == 200, response.text
    assert settings.SESSION_COOKIE_NAME in client.cookies
    return client


@pytest.fixture
def rates() -> RateConfig:
    return RateConfig(
        standard_rate=Decimal("0.15"),
        promo_rate=Decimal("0.10"),
        promo_cutoff=settings.PROMO_CUTOFF,
        line_tolerance=Decimal("0.01"),
        order_tolerance=Decimal("0.05"),
    )


@pytest.fixture
def upload_files() -> dict[str, tuple[str, bytes, str]]:
    """The three required multipart fields, built from the sample workbooks.

    The same workbook goes into ``portal`` and ``promo``: it holds both sheets,
    and each field detects its own.
    """
    xlsx = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    workings = WORKINGS.read_bytes()
    return {
        "portal": ("workings.xlsx", workings, xlsx),
        "promo": ("workings.xlsx", workings, xlsx),
        "vendor": ("vendor_statement.xlsx", VENDOR_STATEMENT.read_bytes(), xlsx),
    }


@pytest.fixture
def travel_rates() -> TravelRateConfig:
    return TravelRateConfig(
        hotel_rate=Decimal("0.05"),
        flight_rate=Decimal("0.03"),
        vat_rate=Decimal("0.15"),
        base_tolerance=Decimal("0.05"),
        commission_tolerance=Decimal("0.01"),
        rate_tolerance=Decimal("0.0005"),
    )


@pytest.fixture
def travel_upload_files() -> dict[str, tuple[str, bytes, str]]:
    """The two required travel fields plus the optional wallet history."""
    xlsx = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return {
        "portal": ("portal.xlsx", TRAVEL_PORTAL.read_bytes(), xlsx),
        "vendor": ("claim_13.xlsx", TRAVEL_CLAIMS["13"].read_bytes(), xlsx),
        "wallet": ("wallet_history.xlsx", TRAVEL_WALLET.read_bytes(), xlsx),
    }


FINANCE_HEADERS = (
    "Spin Event ID",
    "User ID",
    "Order ID",
    "Order Type",
    "Order Amount (SAR)",
    "Reward Amount (SAR)",
    "Reward Code",
    "Reward Type",
    "Outcome",
    "Spin Earned At (KSA)",
    "Reward Earned At (KSA)",
    "AOV Tier ID",
    "AOV Tier Name",
    "Vertical ID",
    "Vertical Name",
)


def write_finance_report(path: Path, order_ids: list[int], vertical: str) -> Path:
    """A spin report as the client sends it: a summary block above the table."""
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet.title = "CASHBACK"
    sheet["A1"] = "Total spins"
    sheet["B1"] = len(order_ids)
    for col, header in enumerate(FINANCE_HEADERS, start=1):
        sheet.cell(14, col, header)
    for offset, order_id in enumerate(order_ids):
        row = 15 + offset
        sheet.cell(row, 1, 61167 + offset)
        sheet.cell(row, 3, order_id)
        sheet.cell(row, 6, 4.13)
        sheet.cell(row, 7, "PERCENT_10")
        sheet.cell(row, 8, 0.1)
        sheet.cell(row, 9, "won")
        sheet.cell(row, 15, vertical)
    book.save(path)
    return path
