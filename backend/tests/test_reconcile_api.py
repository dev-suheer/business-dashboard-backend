import zipfile
from decimal import Decimal
from io import BytesIO
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.core.config import settings

RECONCILE = f"{settings.API_V1_STR}/reconcile"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def test_end_to_end_json(
    auth_client: TestClient, upload_files: dict[str, tuple[str, bytes, str]]
) -> None:
    response = auth_client.post(RECONCILE, files=upload_files)
    assert response.status_code == 200, response.text
    body = response.json()

    assert Decimal(body["summary"]["net_payable"]) == Decimal("13876.31")
    assert body["summary"]["orders"] == 37
    assert body["period_start"] == "2026-07-22"
    # The manifest makes a misdetection visible instead of silent.
    detected = {m["role"]: m["sheets"] for m in body["detected"]}
    assert detected["portal"] == ["E-SHop 1 Vendor "]
    assert detected["promo"] == ["Discount Base Sheet"]


def test_end_to_end_xlsx(
    auth_client: TestClient, upload_files: dict[str, tuple[str, bytes, str]]
) -> None:
    response = auth_client.post(
        RECONCILE, files=upload_files, params={"format": "xlsx"}
    )
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == XLSX
    assert (
        "reconciliation-eshop-2026-07-22-to-2026-08-25.xlsx"
        in (response.headers["content-disposition"])
    )
    with zipfile.ZipFile(BytesIO(response.content)) as book:
        assert "xl/workbook.xml" in book.namelist()


def test_the_wallet_file_is_optional(
    auth_client: TestClient, upload_files: dict[str, tuple[str, bytes, str]]
) -> None:
    from tests.conftest import TRAVEL_WALLET

    assert auth_client.post(RECONCILE, files=upload_files).json()["wallet"] is None
    upload_files["wallet"] = (
        "wallet_history.xlsx",
        TRAVEL_WALLET.read_bytes(),
        XLSX,
    )
    response = auth_client.post(RECONCILE, files=upload_files)
    assert response.status_code == 200, response.text
    body = response.json()
    assert Decimal(body["summary"]["net_payable"]) == Decimal("13876.31")
    states = {c["booking_id"]: c["state"] for c in body["wallet"]["checks"]}
    assert states["SG-SH-08986"] == "matched"
    report = auth_client.post(RECONCILE, files=upload_files, params={"format": "xlsx"})
    assert report.status_code == 200, report.text
    with zipfile.ZipFile(BytesIO(report.content)) as book:
        assert 'name="Wallet"' in book.read("xl/workbook.xml").decode()


def test_a_finance_report_is_optional_and_copied(
    auth_client: TestClient,
    upload_files: dict[str, tuple[str, bytes, str]],
    tmp_path: Path,
) -> None:
    from tests.conftest import write_finance_report

    finance = write_finance_report(tmp_path / "finance.xlsx", [8986], "eshop")
    upload_files["finance"] = ("finance.xlsx", finance.read_bytes(), XLSX)
    response = auth_client.post(
        RECONCILE, files=upload_files, params={"format": "xlsx"}
    )
    assert response.status_code == 200, response.text
    book = load_workbook(BytesIO(response.content), read_only=True)
    assert book.sheetnames[-1] == "Finance Report"
    assert "B2C Invoices ERP" in book.sheetnames
    book.close()


def test_a_finance_file_without_the_spin_table_is_rejected(
    auth_client: TestClient, upload_files: dict[str, tuple[str, bytes, str]]
) -> None:
    upload_files["finance"] = ("finance.xlsx", upload_files["vendor"][1], XLSX)
    response = auth_client.post(RECONCILE, files=upload_files)
    assert response.status_code == 422
    assert "finance" in response.json()["detail"]


def test_a_run_level_rate_override_changes_the_result(
    auth_client: TestClient, upload_files: dict[str, tuple[str, bytes, str]]
) -> None:
    """A back-dated recalculation is one request, not a redeploy."""
    baseline = auth_client.post(RECONCILE, files=upload_files).json()
    shifted = auth_client.post(
        RECONCILE, files=upload_files, data={"promo_cutoff": "2026-07-01"}
    ).json()
    # With the cutoff moved before every order, nothing qualifies for 10%.
    assert Decimal(shifted["summary"]["promo_base"]) == Decimal("0")
    assert Decimal(baseline["summary"]["promo_base"]) == Decimal("4473.10")
    assert shifted["rates"]["promo_cutoff"] == "2026-07-01"


def test_explicit_period_overrides_the_derived_one(
    auth_client: TestClient, upload_files: dict[str, tuple[str, bytes, str]]
) -> None:
    response = auth_client.post(
        RECONCILE,
        files=upload_files,
        data={"period_start": "2026-01-01", "period_end": "2026-12-31"},
    )
    body = response.json()
    assert body["period_start"] == "2026-01-01"
    # A window that swallows the whole statement turns the informational
    # prior/later rows into real breaks, which is the point of the split.
    counts = body["summary"]["exception_counts"]
    assert counts.get("PRIOR_PERIOD", 0) == 0
    assert counts.get("MISSING_FROM_PORTAL", 0) == 21


def test_wrong_file_in_a_field_is_a_422_naming_the_field(
    auth_client: TestClient, upload_files: dict[str, tuple[str, bytes, str]]
) -> None:
    swapped = dict(upload_files)
    swapped["portal"] = upload_files["vendor"]
    response = auth_client.post(RECONCILE, files=swapped)
    assert response.status_code == 422
    detail = response.json()["detail"]
    assert "portal" in detail and "vendor_statement.xlsx" in detail


def test_a_damaged_workbook_is_a_422_naming_the_field(
    auth_client: TestClient, upload_files: dict[str, tuple[str, bytes, str]]
) -> None:
    upload_files["portal"] = ("portal.xlsx", b"not a workbook", XLSX)
    response = auth_client.post(RECONCILE, files=upload_files)
    assert response.status_code == 422
    assert "'portal' is not a readable Excel workbook" in response.json()["detail"]


def test_a_non_excel_upload_is_rejected(
    auth_client: TestClient, upload_files: dict[str, tuple[str, bytes, str]]
) -> None:
    bad = dict(upload_files)
    bad["portal"] = ("notes.csv", b"a,b,c", "text/csv")
    response = auth_client.post(RECONCILE, files=bad)
    assert response.status_code == 415


def test_an_unknown_category_is_rejected(
    auth_client: TestClient, upload_files: dict[str, tuple[str, bytes, str]]
) -> None:
    response = auth_client.post(
        RECONCILE, files=upload_files, data={"category": "groceries"}
    )
    assert response.status_code == 422
    assert "eshop" in response.json()["detail"]


def test_an_oversized_upload_is_rejected_by_the_cap(
    auth_client: TestClient,
    upload_files: dict[str, tuple[str, bytes, str]],
    monkeypatch,
) -> None:
    monkeypatch.setattr(settings, "MAX_UPLOAD_BYTES", 1024)
    response = auth_client.post(RECONCILE, files=upload_files)
    assert response.status_code == 413


def test_missing_files_is_a_validation_error(auth_client: TestClient) -> None:
    assert auth_client.post(RECONCILE).status_code == 422


def test_temporary_files_do_not_accumulate(
    auth_client: TestClient, upload_files: dict[str, tuple[str, bytes, str]]
) -> None:
    """Both the spooled uploads and the generated report are cleaned up."""
    import tempfile
    from pathlib import Path

    tmp = Path(tempfile.gettempdir())
    before = set(tmp.glob("recon-*"))
    auth_client.post(RECONCILE, files=upload_files)
    auth_client.post(RECONCILE, files=upload_files, params={"format": "xlsx"})
    assert set(tmp.glob("recon-*")) == before


def test_a_failed_report_build_leaves_nothing_behind(
    auth_client: TestClient,
    upload_files: dict[str, tuple[str, bytes, str]],
    monkeypatch,
) -> None:
    """The xlsx path owns two temp paths; a mid-build failure must free both."""
    import tempfile
    from pathlib import Path

    from app.api.routes import reconcile as route

    def boom(*_args: object, **_kwargs: object) -> None:
        raise RuntimeError("disk full")

    monkeypatch.setattr(route, "build_report", boom)
    tmp = Path(tempfile.gettempdir())
    before = set(tmp.glob("recon-*"))
    with pytest.raises(RuntimeError, match="disk full"):
        auth_client.post(RECONCILE, files=upload_files, params={"format": "xlsx"})
    assert set(tmp.glob("recon-*")) == before
