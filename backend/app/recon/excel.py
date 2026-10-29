"""Constant-memory workbook reading.

openpyxl's read-only mode streams one row at a time, which is what lets a
multi-hundred-megabyte upload be processed without ever being fully resident.
Two rules follow from that and are enforced here:

* a worksheet is iterated, never indexed or measured;
* headers are located by text, never by column position — the source sheets
  shift every column by one between pages.
"""

import zlib
from collections.abc import Iterator, Mapping, Sequence
from contextlib import contextmanager
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.recon.cleaning import normalize_header

# What openpyxl raises for a file that is not a readable workbook: not a zip,
# a zip with no workbook in it, or a damaged part. A missing part is a
# KeyError, which only counts while the workbook is being opened; once it is
# open, a KeyError is a fault in the code reading it, not in the file.
_UNREADABLE_ON_OPEN = (BadZipFile, KeyError, ParseError, InvalidFileException)
# A read-only sheet is decompressed and parsed as it is iterated.
_UNREADABLE_ON_READ = (BadZipFile, zlib.error, ParseError)

# How many rows to inspect while hunting for the header. The vendor ledger puts
# it on row 4 behind account metadata; nothing legitimate hides deeper.
HEADER_SCAN_LIMIT = 30


class SheetFormatError(ValueError):
    """A sheet does not carry the headers its role requires."""


@dataclass(frozen=True)
class ColumnSpec:
    """One logical field and the header texts that may introduce it."""

    field: str
    aliases: tuple[str, ...]
    required: bool = False

    @property
    def keys(self) -> tuple[str, ...]:
        return tuple(normalize_header(alias) for alias in self.aliases)


@dataclass(frozen=True)
class SheetSchema:
    """The shape of one kind of sheet, plus how to recognise it."""

    name: str
    columns: tuple[ColumnSpec, ...]
    # Every one of these normalised headers must be present for a sheet to be
    # accepted as this kind.
    fingerprint: tuple[str, ...] = field(default=())
    # Headers whose presence disqualifies a sheet. The source workbook's legacy
    # recon tab carries both the portal columns and the promo columns bolted
    # together, so it satisfies either fingerprint; a disqualifier is what keeps
    # it from being mistaken for an input.
    exclude: tuple[str, ...] = field(default=())

    @property
    def required_fields(self) -> tuple[str, ...]:
        return tuple(spec.field for spec in self.columns if spec.required)

    def fingerprint_keys(self) -> tuple[str, ...]:
        return tuple(normalize_header(header) for header in self.fingerprint)

    def exclude_keys(self) -> tuple[str, ...]:
        return tuple(normalize_header(header) for header in self.exclude)


@dataclass(frozen=True)
class HeaderMatch:
    """Where a sheet's header sits and which column index feeds each field."""

    row: int
    columns: Mapping[str, int]

    def missing(self, schema: SheetSchema) -> tuple[str, ...]:
        return tuple(f for f in schema.required_fields if f not in self.columns)


@contextmanager
def open_workbook(path: Path) -> Iterator[Workbook]:
    """Open a workbook for streaming reads, always closing it again.

    ``read_only`` streams rows; ``data_only`` returns the cached result of a
    formula rather than its text, which matters because the source workbooks are
    full of formulas.

    A damaged file is a ``SheetFormatError``, so the upload is refused with a
    reason instead of failing the request. Uploads are spooled as
    ``<field>.xlsx``, so the file's stem names the field in that reason.
    """
    try:
        workbook = load_workbook(filename=path, read_only=True, data_only=True)
    except _UNREADABLE_ON_OPEN as exc:
        raise _unreadable(path) from exc
    try:
        yield workbook
    except _UNREADABLE_ON_READ as exc:
        raise _unreadable(path) from exc
    finally:
        # Read-only workbooks hold an open zip handle until closed.
        workbook.close()


def _unreadable(path: Path) -> SheetFormatError:
    return SheetFormatError(
        f"'{path.stem}' is not a readable Excel workbook. Open it in Excel, "
        "save it as .xlsx and upload it again."
    )


def _row_headers(row: Sequence[Any]) -> dict[str, int]:
    headers: dict[str, int] = {}
    for index, value in enumerate(row):
        key = normalize_header(value)
        # First occurrence wins: several sheets repeat a header further right
        # (the portal export carries "Order Number" twice).
        if key and key not in headers:
            headers[key] = index
    return headers


def _match_columns(
    headers: Mapping[str, int], schema: SheetSchema
) -> dict[str, int] | None:
    if not all(key in headers for key in schema.fingerprint_keys()):
        return None
    if any(key in headers for key in schema.exclude_keys()):
        return None
    columns: dict[str, int] = {}
    for spec in schema.columns:
        for key in spec.keys:
            if key in headers:
                columns[spec.field] = headers[key]
                break
    return columns


def find_header(worksheet: Worksheet, schema: SheetSchema) -> HeaderMatch | None:
    """Locate the header row, or return ``None`` if this is not such a sheet."""
    for row_number, row in enumerate(
        worksheet.iter_rows(max_row=HEADER_SCAN_LIMIT, values_only=True), start=1
    ):
        columns = _match_columns(_row_headers(row), schema)
        if columns is not None:
            return HeaderMatch(row=row_number, columns=columns)
    return None


def iter_records(
    worksheet: Worksheet,
    schema: SheetSchema,
    *,
    max_rows: int,
) -> Iterator[dict[str, Any]]:
    """Stream a sheet's data rows as ``{field: raw value}`` dicts.

    Single pass: the header is detected on the same iterator that goes on to
    yield the data, so the sheet is read exactly once. Values are returned raw —
    coercion belongs to the category parser, which knows what each field means.
    """
    match: HeaderMatch | None = None
    columns: Mapping[str, int] = {}
    width = 0
    seen = 0

    for row_number, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
        if match is None:
            if row_number > HEADER_SCAN_LIMIT:
                break
            found = _match_columns(_row_headers(row), schema)
            if found is None:
                continue
            match = HeaderMatch(row=row_number, columns=found)
            missing = match.missing(schema)
            if missing:
                raise SheetFormatError(
                    f"{schema.name} sheet {worksheet.title!r} is missing required "
                    f"column(s): {', '.join(missing)}"
                )
            columns = match.columns
            width = max(columns.values()) + 1
            continue

        seen += 1
        if seen > max_rows:
            raise SheetFormatError(
                f"Sheet {worksheet.title!r} exceeds the {max_rows:,}-row limit"
            )
        if len(row) < width:
            # Trailing cells are omitted entirely on short rows.
            row = tuple(row) + (None,) * (width - len(row))
        yield {field_name: row[index] for field_name, index in columns.items()}

    if match is None:
        raise SheetFormatError(
            f"No {schema.name} header found in sheet {worksheet.title!r}"
        )


def detect_sheet(workbook: Workbook, schema: SheetSchema) -> str | None:
    """Return the name of the first sheet matching ``schema``, if any.

    Callers upload whole workbooks, so a single file may hold the portal export,
    the promo catalogue, a legacy recon tab and another category's data. Peeking
    at each sheet's header block is cheap next to reading its body.
    """
    for name in workbook.sheetnames:
        if find_header(workbook[name], schema) is not None:
            return str(name)
    return None


def detect_sheets(workbook: Workbook, schema: SheetSchema) -> list[str]:
    """Return every sheet matching ``schema``.

    The vendor ledger needs this: the PDF converter emits overlapping copies of
    the same transactions across several sheets, all of which must be read and
    then deduplicated.
    """
    return [
        str(name)
        for name in workbook.sheetnames
        if find_header(workbook[name], schema) is not None
    ]
